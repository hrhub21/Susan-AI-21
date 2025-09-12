#!/usr/bin/env python3
"""
Simplified Sales Data Extractor for Susan AI
===========================================

Extracts and processes sales training materials quickly and efficiently
for integration with Susan AI's roof damage training pipeline.
"""

import os
import sys
import json
import logging
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Any, Optional
import hashlib
import re

# Document processing
try:
    import PyPDF2
    from pptx import Presentation
    from docx import Document
    import openpyxl
    from PIL import Image
except ImportError as e:
    print(f"Missing required package: {e}")
    print("Please install with: pip3 install PyPDF2 python-docx python-pptx openpyxl pillow")
    sys.exit(1)

# Set up logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class SimpleSalesExtractor:
    """Simplified sales data extractor"""
    
    def __init__(self):
        self.output_dir = Path("./sales_training_data")
        self.output_dir.mkdir(exist_ok=True)
        
        self.processed_data_dir = self.output_dir / "processed"
        self.processed_data_dir.mkdir(exist_ok=True)
        
        # Statistics
        self.stats = {
            "files_processed": 0,
            "total_content_extracted": 0,
            "sales_scripts_found": 0,
            "damage_terms_found": 0,
            "errors": 0
        }
    
    def extract_all_sales_data(self) -> Dict[str, Any]:
        """Extract all sales training data"""
        logger.info("Starting sales data extraction")
        
        target_paths = [
            "/Users/a21/Desktop/Sales Rep Resources",
            "/Users/a21/Desktop/Sales Reps",
            "/Users/a21/Desktop/Roof-ER Sales Training.pptx"
        ]
        
        all_extracted_data = {
            "sales_scripts": [],
            "damage_terminology": [],
            "customer_interactions": [],
            "process_workflows": [],
            "visual_materials": [],
            "extraction_stats": {}
        }
        
        for target_path in target_paths:
            logger.info(f"Processing: {target_path}")
            
            if os.path.isfile(target_path):
                file_data = self._process_single_file(target_path)
                if file_data:
                    self._categorize_and_add_data(file_data, all_extracted_data)
            elif os.path.isdir(target_path):
                dir_data = self._process_directory(target_path)
                for file_data in dir_data:
                    if file_data:
                        self._categorize_and_add_data(file_data, all_extracted_data)
            else:
                logger.warning(f"Path not found: {target_path}")
        
        # Save extracted data
        self._save_extracted_data(all_extracted_data)
        
        # Update stats
        all_extracted_data["extraction_stats"] = self.stats
        
        logger.info("Sales data extraction completed")
        logger.info(f"Files processed: {self.stats['files_processed']}")
        logger.info(f"Sales scripts found: {self.stats['sales_scripts_found']}")
        logger.info(f"Damage terms found: {self.stats['damage_terms_found']}")
        
        return all_extracted_data
    
    def _process_directory(self, directory_path: str) -> List[Dict[str, Any]]:
        """Process all files in a directory"""
        extracted_data = []
        
        for root, dirs, files in os.walk(directory_path):
            for file_name in files:
                file_path = os.path.join(root, file_name)
                
                if self._should_process_file(file_path):
                    file_data = self._process_single_file(file_path)
                    if file_data:
                        extracted_data.append(file_data)
        
        return extracted_data
    
    def _process_single_file(self, file_path: str) -> Optional[Dict[str, Any]]:
        """Process a single file"""
        try:
            file_ext = Path(file_path).suffix.lower()
            
            # Extract content based on file type
            content = ""
            if file_ext == ".pdf":
                content = self._extract_pdf(file_path)
            elif file_ext == ".docx":
                content = self._extract_docx(file_path)
            elif file_ext == ".pptx":
                content = self._extract_pptx(file_path)
            elif file_ext == ".xlsx":
                content = self._extract_xlsx(file_path)
            elif file_ext in [".jpg", ".jpeg", ".png"]:
                content = f"[Image file: {os.path.basename(file_path)}]"
            else:
                return None
            
            if not content or len(content.strip()) < 50:
                return None
            
            # Analyze content
            file_data = {
                "file_path": file_path,
                "file_name": os.path.basename(file_path),
                "file_type": file_ext,
                "content": content,
                "word_count": len(content.split()),
                "category": self._categorize_content(content, file_path),
                "sales_patterns": self._identify_sales_patterns(content),
                "damage_terms": self._extract_damage_terms(content),
                "customer_objections": self._find_customer_objections(content),
                "process_steps": self._extract_process_steps(content)
            }
            
            self.stats["files_processed"] += 1
            self.stats["total_content_extracted"] += len(content)
            
            if file_data["sales_patterns"]:
                self.stats["sales_scripts_found"] += 1
            
            if file_data["damage_terms"]:
                self.stats["damage_terms_found"] += len(file_data["damage_terms"])
            
            return file_data
            
        except Exception as e:
            logger.error(f"Error processing {file_path}: {str(e)}")
            self.stats["errors"] += 1
            return None
    
    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF"""
        try:
            content = []
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    text = page.extract_text()
                    if text.strip():
                        content.append(text)
            return "\n".join(content)
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            return ""
    
    def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX"""
        try:
            doc = Document(file_path)
            content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        content.append(" | ".join(row_text))
            
            return "\n".join(content)
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            return ""
    
    def _extract_pptx(self, file_path: str) -> str:
        """Extract text from PowerPoint"""
        try:
            presentation = Presentation(file_path)
            content = []
            
            for slide_num, slide in enumerate(presentation.slides, 1):
                slide_content = [f"[Slide {slide_num}]"]
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text)
                
                if len(slide_content) > 1:
                    content.append("\n".join(slide_content))
            
            return "\n\n".join(content)
        except Exception as e:
            logger.error(f"PPTX extraction error: {e}")
            return ""
    
    def _extract_xlsx(self, file_path: str) -> str:
        """Extract text from Excel"""
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            content = []
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_content = [f"[Sheet: {sheet_name}]"]
                
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) if cell is not None else "" for cell in row]
                    if any(cell.strip() for cell in row_data if cell):
                        sheet_content.append(" | ".join(row_data))
                
                if len(sheet_content) > 1:
                    content.append("\n".join(sheet_content))
            
            return "\n\n".join(content)
        except Exception as e:
            logger.error(f"XLSX extraction error: {e}")
            return ""
    
    def _categorize_content(self, content: str, file_path: str) -> str:
        """Categorize content based on content and path"""
        path_lower = file_path.lower()
        content_lower = content.lower()
        
        if "script" in path_lower or any(word in content_lower for word in ["hello", "good morning", "pitch"]):
            return "sales_script"
        elif "template" in path_lower or "email" in path_lower:
            return "communication_template"
        elif "insurance" in path_lower or "adjuster" in path_lower:
            return "insurance_documentation"
        elif "warranty" in path_lower or "product" in path_lower:
            return "product_information"
        elif "photo" in path_lower or "example" in path_lower:
            return "visual_training"
        elif "process" in path_lower or "operation" in path_lower:
            return "operational_procedure"
        else:
            return "general_content"
    
    def _identify_sales_patterns(self, content: str) -> List[str]:
        """Identify sales patterns in content"""
        patterns = []
        content_lower = content.lower()
        
        # Greeting patterns
        if any(word in content_lower for word in ["hello", "good morning", "introduce"]):
            patterns.append("greeting")
        
        # Value proposition
        if any(word in content_lower for word in ["save money", "protect", "free"]):
            patterns.append("value_proposition")
        
        # Urgency
        if any(word in content_lower for word in ["limited time", "act now", "deadline"]):
            patterns.append("urgency")
        
        # Objection handling
        if any(word in content_lower for word in ["understand", "concern", "however"]):
            patterns.append("objection_handling")
        
        # Closing
        if any(word in content_lower for word in ["sign", "agreement", "move forward"]):
            patterns.append("closing")
        
        return patterns
    
    def _extract_damage_terms(self, content: str) -> List[str]:
        """Extract damage-related terms"""
        damage_terms = [
            "hail", "wind", "storm", "damage", "shingle", "roof", "leak", 
            "repair", "replace", "granule", "impact", "missing", "torn",
            "uplift", "cracking", "curling", "wear", "deterioration"
        ]
        
        found_terms = []
        content_lower = content.lower()
        
        for term in damage_terms:
            if term in content_lower:
                found_terms.append(term)
        
        return found_terms
    
    def _find_customer_objections(self, content: str) -> List[str]:
        """Find customer objection patterns"""
        objections = []
        content_lower = content.lower()
        
        objection_patterns = [
            "too expensive", "can't afford", "need to think", "get other quotes",
            "not interested", "insurance won't cover", "deductible too high"
        ]
        
        for objection in objection_patterns:
            if objection in content_lower:
                objections.append(objection)
        
        return objections
    
    def _extract_process_steps(self, content: str) -> List[str]:
        """Extract process steps"""
        steps = []
        
        # Look for numbered steps
        step_pattern = r'(?:step\s*)?(\d+)[.\)\:]?\s*([^.\n]+)'
        matches = re.findall(step_pattern, content, re.IGNORECASE)
        
        for step_num, step_text in matches:
            clean_step = step_text.strip()
            if len(clean_step) > 10:
                steps.append(f"Step {step_num}: {clean_step}")
        
        return steps
    
    def _categorize_and_add_data(self, file_data: Dict[str, Any], all_data: Dict[str, Any]):
        """Categorize and add file data to appropriate collections"""
        category = file_data["category"]
        
        if category == "sales_script":
            all_data["sales_scripts"].append(file_data)
        elif category in ["insurance_documentation", "product_information"]:
            all_data["damage_terminology"].extend(file_data["damage_terms"])
        elif category == "communication_template":
            all_data["customer_interactions"].append(file_data)
        elif category == "operational_procedure":
            all_data["process_workflows"].append(file_data)
        elif category == "visual_training":
            all_data["visual_materials"].append(file_data)
    
    def _save_extracted_data(self, data: Dict[str, Any]):
        """Save extracted data to files"""
        
        # Save full data as JSON
        json_path = self.processed_data_dir / "extracted_sales_data.json"
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        
        # Save individual datasets
        for dataset_name, dataset in data.items():
            if dataset_name != "extraction_stats" and dataset:
                file_path = self.processed_data_dir / f"{dataset_name}.json"
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(dataset, f, indent=2, ensure_ascii=False)
        
        # Create summary report
        summary = {
            "extraction_date": datetime.now().isoformat(),
            "statistics": self.stats,
            "datasets_created": {
                name: len(data) if isinstance(data, list) else len(data) if isinstance(data, dict) else 0
                for name, data in data.items() if name != "extraction_stats"
            }
        }
        
        summary_path = self.processed_data_dir / "extraction_summary.json"
        with open(summary_path, 'w', encoding='utf-8') as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)
        
        logger.info(f"Data saved to: {self.processed_data_dir}")
    
    def _should_process_file(self, file_path: str) -> bool:
        """Check if file should be processed"""
        file_ext = Path(file_path).suffix.lower()
        supported_formats = [".pdf", ".docx", ".pptx", ".xlsx", ".jpg", ".jpeg", ".png"]
        
        if file_ext not in supported_formats:
            return False
        
        # Check file size (skip very large files)
        try:
            file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
            if file_size_mb > 50:  # Skip files larger than 50MB
                return False
        except OSError:
            return False
        
        # Skip system files
        file_name = os.path.basename(file_path).lower()
        skip_patterns = ['.ds_store', 'thumbs.db', '~$']
        if any(pattern in file_name for pattern in skip_patterns):
            return False
        
        return True

def main():
    """Main function"""
    print("="*60)
    print("SUSAN AI SALES TRAINING DATA EXTRACTOR")
    print("="*60)
    
    extractor = SimpleSalesExtractor()
    
    try:
        results = extractor.extract_all_sales_data()
        
        print(f"\nExtraction completed successfully!")
        print(f"Files processed: {results['extraction_stats']['files_processed']}")
        print(f"Sales scripts found: {results['extraction_stats']['sales_scripts_found']}")
        print(f"Damage terms found: {results['extraction_stats']['damage_terms_found']}")
        print(f"Errors encountered: {results['extraction_stats']['errors']}")
        
        print(f"\nDatasets created:")
        for name, data in results.items():
            if name != "extraction_stats" and data:
                count = len(data) if isinstance(data, list) else len(data) if isinstance(data, dict) else 0
                print(f"- {name}: {count} items")
        
        print(f"\nData saved to: ./sales_training_data/processed/")
        print("="*60)
        
        return True
        
    except Exception as e:
        print(f"Extraction failed: {str(e)}")
        return False

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)